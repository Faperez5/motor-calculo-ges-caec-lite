"""LOAD: every write to the database, the text reports, and the Excel template.

Results are versioned: a new version marks the previous one vigente = 0 and
never overwrites it. Facts (prestaciones) are append-only (insertar_prestaciones). Rules only change
through cargar_regla_ges (plan D10). Beneficiaries can be added, never changed (insertar_beneficiarios).
"""

from datetime import date, datetime

import openpyxl
from openpyxl.styles import Font

from src import calculo, extractor
from src.calculo import Resultado


def iniciar_ejecucion(conn, comando: str, usuario: str, version_codigo: str, iniciado_en: datetime) -> int:
    """Log a command as started, committed at once so a crash still leaves a trace."""
    with conn:
        cur = conn.execute(
            "INSERT INTO ejecuciones (comando, iniciado_en, usuario, version_codigo, estado) "
            "VALUES (?, ?, ?, ?, 'en_curso')",
            (comando, iniciado_en.isoformat(timespec="seconds"), usuario, version_codigo),
        )
    return cur.lastrowid


def terminar_ejecucion(conn, ejecucion_id: int, estado: str, mensaje: str | None, terminado_en: datetime) -> None:
    with conn:
        conn.execute(
            "UPDATE ejecuciones SET estado = ?, mensaje = ?, terminado_en = ? WHERE id = ?",
            (estado, mensaje, terminado_en.isoformat(timespec="seconds"), ejecucion_id),
        )


def guardar_resultado(conn, resultado: Resultado, motivo: str, calculado_en: datetime,
                      conocido_hasta: datetime, version_codigo: str, advertencias=(),
                      ejecucion_id: int | None = None) -> int:
    """Save `resultado` as the new current version. Returns its id.

    One transaction: the previous current version (if any) is marked vigente = 0
    and the new one gets version_calculo = previous + 1. If anything fails,
    nothing changes, including the previous version's vigente flag.
    `sin_cambios` is refused: an unchanged result must not create a version.
    So is a `conocido_hasta` earlier than the current version's: a result that
    knew less must not replace one that knew more (look at the past with
    `calcular --simular` instead).
    """
    if motivo == "sin_cambios":
        raise ValueError("sin_cambios: nothing to save")
    clave = (resultado.beneficiario_id, resultado.tipo)
    with conn:
        vigente = conn.execute(
            "SELECT conocido_hasta FROM calculos_resultado WHERE beneficiario_id = ? AND tipo = ? AND vigente = 1", clave
        ).fetchone()
        if vigente and conocido_hasta.isoformat(timespec="seconds") < vigente[0]:
            raise ValueError(f"the current result already knows data up to {vigente[0]}; saving one that only "
                             f"knows up to {conocido_hasta.isoformat(timespec='seconds')} would replace it with "
                             "less. Use --simular to look at the past")
        ultima = conn.execute(
            "SELECT max(version_calculo) FROM calculos_resultado WHERE beneficiario_id = ? AND tipo = ?", clave
        ).fetchone()[0]
        conn.execute(
            "UPDATE calculos_resultado SET vigente = 0 WHERE beneficiario_id = ? AND tipo = ? AND vigente = 1", clave
        )
        cur = conn.execute(
            """
            INSERT INTO calculos_resultado
                (beneficiario_id, tipo, version_calculo, regla_version_id,
                 ventana_inicio, ventana_fin, acumulado, deducible_aplicado,
                 fecha_cruce, calculado_en, conocido_hasta, version_codigo,
                 advertencias, ejecucion_id, motivo_recalculo, vigente)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 1)
            """,
            (
                resultado.beneficiario_id,
                resultado.tipo,
                (ultima or 0) + 1,
                resultado.regla_version_id,
                resultado.ventana_inicio.isoformat(),
                resultado.ventana_fin.isoformat(),
                resultado.acumulado,
                resultado.deducible_aplicado,
                resultado.fecha_cruce.isoformat() if resultado.fecha_cruce else None,
                calculado_en.isoformat(timespec="seconds"),
                conocido_hasta.isoformat(timespec="seconds"),
                version_codigo,
                "\n".join(advertencias) or None,
                ejecucion_id,
                motivo,
            ),
        )
        calculo_id = cur.lastrowid
        conn.executemany(
            """
            INSERT INTO calculos_detalle
                (calculo_id, prestacion_id, copago, paga_beneficiario, cubre_seguro, acumulado_hasta_aqui)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            [
                (calculo_id, f.prestacion_id, f.copago, f.paga_beneficiario, f.cubre_seguro, f.acumulado)
                for f in resultado.filas
            ],
        )
    return calculo_id


def insertar_beneficiarios(conn, filas: list[dict]) -> tuple[list[int], list[int], list[tuple[dict, str]]]:
    """Add new beneficiaries. Returns (new ids, ids already stored unchanged, rejected).

    One transaction. A beneficiary already stored with the same data is left
    alone (loading the same file twice is safe); one stored with different
    data is rejected: changing a beneficiary is out of scope.
    """
    nuevos, iguales, rechazadas = [], [], []
    with conn:
        for f in filas:
            valores = (f["cotizacion_pactada_uf"],
                       f["fecha_derivacion_caec"].isoformat() if f["fecha_derivacion_caec"] else None)
            actual = conn.execute(
                "SELECT cotizacion_pactada_uf, fecha_derivacion_caec FROM beneficiarios WHERE beneficiario_id = ?",
                (f["beneficiario_id"],),
            ).fetchone()
            if actual is None:
                conn.execute(
                    "INSERT INTO beneficiarios (beneficiario_id, cotizacion_pactada_uf, fecha_derivacion_caec) "
                    "VALUES (?, ?, ?)", (f["beneficiario_id"], *valores))
                nuevos.append(f["beneficiario_id"])
            elif tuple(actual) == valores:
                iguales.append(f["beneficiario_id"])
            else:
                rechazadas.append((f, f"beneficiary {f['beneficiario_id']} already exists with different data "
                                      f"(cotización {actual[0]} UF, CAEC referral {actual[1] or 'none'}); "
                                      "changing a beneficiary isn't supported"))
    return nuevos, iguales, rechazadas


def insertar_prestaciones(conn, filas: list[dict], recibido_en: datetime,
                          ejecucion_id: int | None = None) -> tuple[list[int], list[tuple[dict, str]]]:
    """Append ingested prestaciones (append-only: never updates). Returns (new ids, rejected).

    One transaction. A row is rejected, not inserted, when its beneficiary
    doesn't exist or when an identical prestación is already stored (same
    beneficiary, date, code, networks and amounts). That catches a file loaded
    twice and overlapping exports; it also means two truly identical copagos
    on the same day need a manual check.
    """
    nuevas, rechazadas = [], []
    with conn:
        for f in filas:
            if conn.execute("SELECT 1 FROM beneficiarios WHERE beneficiario_id = ?",
                            (f["beneficiario_id"],)).fetchone() is None:
                rechazadas.append((f, f"beneficiary {f['beneficiario_id']} does not exist"))
                continue
            clave = (f["beneficiario_id"], f["fecha_prestacion"].isoformat(), f["codigo_ges"],
                     int(f["en_red_ges"]), int(f["en_red_caec"]), f["monto_arancel"], f["monto_copago"])
            igual = conn.execute(
                """
                SELECT prestacion_id FROM prestaciones
                WHERE beneficiario_id = ? AND fecha_prestacion = ? AND codigo_ges IS ?
                  AND en_red_ges = ? AND en_red_caec = ? AND monto_arancel = ? AND monto_copago = ?
                """,
                clave,
            ).fetchone()
            if igual is not None:
                rechazadas.append((f, f"duplicate of prestación {igual[0]}"))
                continue
            cur = conn.execute(
                """
                INSERT INTO prestaciones
                    (beneficiario_id, fecha_prestacion, codigo_ges, en_red_ges, en_red_caec,
                     monto_arancel, monto_copago, recibido_en, ejecucion_id)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (*clave, recibido_en.isoformat(timespec="seconds"), ejecucion_id),
            )
            nuevas.append(cur.lastrowid)
    return nuevas, rechazadas


def cargar_regla_ges(conn, codigo_ges: str, valid_from: date, tope_cotizaciones: int, tope_uf: float,
                     fuente: str | None, cargado_por: str, cargado_en: datetime,
                     ejecucion_id: int | None = None) -> int:
    """Load a new GES rule version (plan D10). Returns its regla_id.

    One transaction: close the version that contains `valid_from`, insert the
    new one, re-read the whole set and validate it (§8.1). Any problem rolls
    everything back.
    """
    with conn:
        cierres, valid_to = calculo.planear_nueva_version(
            extractor.leer_reglas_ges(conn, codigo_ges), codigo_ges, valid_from)
        for regla_id, hasta in cierres:
            conn.execute("UPDATE parametros_ges SET valid_to = ? WHERE regla_id = ?", (hasta.isoformat(), regla_id))
        cur = conn.execute(
            """
            INSERT INTO parametros_ges
                (codigo_ges, valid_from, valid_to, tope_cotizaciones, tope_uf, fuente, cargado_por, cargado_en,
                 ejecucion_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (codigo_ges, valid_from.isoformat(), valid_to.isoformat() if valid_to else None,
             tope_cotizaciones, tope_uf, fuente, cargado_por, cargado_en.isoformat(timespec="seconds"),
             ejecucion_id),
        )
        problemas = calculo.validar_reglas(extractor.leer_reglas_ges(conn, codigo_ges))
        if problemas:
            raise ValueError("Rule set would be invalid: " + "; ".join(problemas))
    return cur.lastrowid


# ───────────────────────────────────────────────────────────
# Reports: built from a saved result (extractor.leer_resultado),
# never from in-memory values, so they can be regenerated later.
# ───────────────────────────────────────────────────────────

def reporte_texto(r: dict) -> str:
    """Audit trace of a saved result, for the terminal. `id` None = a simulation (reporte_simulacion)."""
    if r["id"] is None:
        cabecera = (f"SIMULATION, NOT SAVED — beneficiary {r['beneficiario_id']}, {r['tipo']}, "
                    f"compared with the current result: {r['motivo_recalculo']}")
    else:
        estado = "current" if r["vigente"] else "superseded"
        cabecera = (f"Result #{r['id']} — beneficiary {r['beneficiario_id']}, {r['tipo']}, "
                    f"version {r['version_calculo']} ({estado}), motivo: {r['motivo_recalculo']}")
    lineas = [
        cabecera,
        f"  rule version    {r['regla_version_id']}",
        f"  window          {r['ventana_inicio']} → {r['ventana_fin']}",
        f"  deductible      {r['deducible_aplicado']:>12,} CLP",
        f"  accumulated     {r['acumulado']:>12,} CLP",
        f"  crossed on      {r['fecha_cruce'] or 'not reached'}",
        f"  calculated at   {r['calculado_en']}   (data known up to {r['conocido_hasta']})",
        f"  code version    {r['version_codigo']}",
        *[f"  WARNING         {a}" for a in (r["advertencias"] or "").splitlines()],
        "",
        f"  {'id':>3}  {'fecha':<10}  {'copago':>11}  {'beneficiary':>11}  {'insurer':>11}  {'accumulated':>11}",
    ]
    for d in r["detalle"]:
        lineas.append(
            f"  {d['prestacion_id']:>3}  {d['fecha_prestacion']:<10}  {d['copago']:>11,}  "
            f"{d['paga_beneficiario']:>11,}  {d['cubre_seguro']:>11,}  {d['acumulado_hasta_aqui']:>11,}"
        )
    return "\n".join(lineas)


def reporte_simulacion(r: Resultado, motivo: str, calculado_en: datetime, conocido_hasta: datetime,
                       version_codigo: str, advertencias=()) -> str:
    """The same trace for a result that was never saved (`calcular --simular`).

    The one report built from in-memory values, on purpose: there is nothing
    stored to rebuild it from.
    """
    return reporte_texto({
        "id": None, "beneficiario_id": r.beneficiario_id, "tipo": r.tipo, "motivo_recalculo": motivo,
        "regla_version_id": r.regla_version_id,
        "ventana_inicio": r.ventana_inicio.isoformat(), "ventana_fin": r.ventana_fin.isoformat(),
        "deducible_aplicado": r.deducible_aplicado, "acumulado": r.acumulado,
        "fecha_cruce": r.fecha_cruce.isoformat() if r.fecha_cruce else None,
        "calculado_en": calculado_en.isoformat(timespec="seconds"),
        "conocido_hasta": conocido_hasta.isoformat(timespec="seconds"),
        "version_codigo": version_codigo, "advertencias": "\n".join(advertencias) or None,
        "detalle": [
            {"prestacion_id": f.prestacion_id, "fecha_prestacion": f.fecha.isoformat(), "copago": f.copago,
             "paga_beneficiario": f.paga_beneficiario, "cubre_seguro": f.cubre_seguro,
             "acumulado_hasta_aqui": f.acumulado}
            for f in r.filas
        ],
    })


def reporte_ficha(ficha: dict) -> str:
    """One beneficiary's data as stored: who, which prestaciones (and when each arrived), current results."""
    b = ficha["beneficiario"]
    visto = f"known up to {ficha['conocido_hasta']}" if ficha["conocido_hasta"] else "everything known now"
    ps = ficha["prestaciones"]
    lineas = [
        f"Beneficiary {b['beneficiario_id']} — cotización pactada {b['cotizacion_pactada_uf']} UF, "
        f"CAEC referral {b['fecha_derivacion_caec'] or 'none'}",
        "",
        f"Prestaciones: {len(ps)} ({visto})",
    ]
    if ps:
        lineas.append(f"  {'id':>3}  {'fecha':<10}  {'received':<19}  {'GES':<6}  {'net GES':<7}  {'net CAEC':<8}  "
                      f"{'arancel':>11}  {'copago':>11}  loaded by")
        for p in ps:
            lineas.append(
                f"  {p['prestacion_id']:>3}  {p['fecha_prestacion']:<10}  {p['recibido_en']:<19}  "
                f"{p['codigo_ges'] or '-':<6}  {'yes' if p['en_red_ges'] else 'no':<7}  "
                f"{'yes' if p['en_red_caec'] else 'no':<8}  {p['monto_arancel']:>11,}  {p['monto_copago']:>11,}  "
                f"{'run #' + str(p['ejecucion_id']) if p['ejecucion_id'] else 'seed'}"
            )
        lineas.append(f"  {'total':<63}  {sum(p['monto_arancel'] for p in ps):>11,}  "
                      f"{sum(p['monto_copago'] for p in ps):>11,}")
    lineas += ["", "Current results"]
    for r in ficha["resultados"]:
        lineas.append(
            f"  {r['tipo']:<4}  result #{r['id']}, version {r['version_calculo']} ({r['motivo_recalculo']}), "
            f"rule {r['regla_version_id']}, window {r['ventana_inicio']} → {r['ventana_fin']}, "
            f"deductible {r['deducible_aplicado']:,} CLP, crossed on {r['fecha_cruce'] or 'not reached'}, "
            f"calculated at {r['calculado_en']}"
        )
    if not ficha["resultados"]:
        lineas.append("  none saved yet: run `calcular`")
    return "\n".join(lineas)


def reporte_historial(filas: list[dict]) -> str:
    """One line per saved version, to see how and why a result changed."""
    if not filas:
        return "no saved results"
    lineas = [f"  {'ver':>3}  {'':1}  {'motivo':<18}  {'rule':>4}  {'window':<23}  {'deductible':>11}  "
              f"{'crossed on':<10}  calculated at"]
    for f in filas:
        lineas.append(
            f"  {f['version_calculo']:>3}  {'*' if f['vigente'] else ' '}  {f['motivo_recalculo']:<18}  "
            f"{f['regla_version_id']:>4}  {f['ventana_inicio']} → {f['ventana_fin']}  {f['deducible_aplicado']:>11,}  "
            f"{f['fecha_cruce'] or '-':<10}  {f['calculado_en']}"
        )
    return "\n".join(lineas) + "\n  (* = current)"


def reporte_ejecuciones(filas: list[dict]) -> str:
    """One line per logged command, newest first."""
    if not filas:
        return "no runs logged"
    lineas = [f"  {'id':>4}  {'started':<19}  {'user':<10}  {'state':<8}  {'code':<8}  command → outcome"]
    for f in filas:
        lineas.append(
            f"  {f['id']:>4}  {f['iniciado_en']:<19}  {f['usuario'][:10]:<10}  {f['estado']:<8}  "
            f"{f['version_codigo'][:7] + ('*' if f['version_codigo'].endswith('-dirty') else ''):<8}  "
            f"{f['comando']}" + (f"  →  {f['mensaje']}" if f["mensaje"] else "")
        )
    return "\n".join(lineas) + "\n  (code* = uncommitted changes)"


# ───────────────────────────────────────────────────────────
# Excel template: a workbook to edit by hand and load with `ingerir --formato A`
# ───────────────────────────────────────────────────────────

# Synthetic examples that fit the seed (beneficiaries 1 and 2 already exist).
EJEMPLO_BENEFICIARIOS = [
    # id, cotización pactada UF, CAEC referral
    (3, 4.0, None),                  # 29 × 4.0 = 116 UF, below the 122 UF tope
    (4, 2.0, None),                  # 58 UF: a small deductible that isn't reached
    (5, 6.5, date(2025, 3, 1)),      # 29 × 6.5 = 188.5 UF, capped at the tope; referred to CAEC
]
EJEMPLO_PRESTACIONES = [
    # beneficiary, date, GES code, GES network, CAEC network, arancel, copago
    (3, date(2025, 6, 12), "GES01", True, False, 4_500_000, 900_000),
    (3, date(2025, 8, 20), "GES01", True, False, 6_000_000, 1_200_000),
    (3, date(2025, 9, 1), None, True, False, 750_000, 150_000),              # not GES: doesn't count
    (3, date(2025, 10, 5), "GES01", True, False, 5_500_000, 1_100_000),
    (3, date(2025, 11, 10), "GES01", False, False, 2_000_000, 400_000),      # outside the GES network: doesn't count
    (3, date(2026, 1, 15), "GES01", True, False, 5_000_000, 1_000_000),
    (3, date(2026, 4, 2), "GES01", True, False, 4_000_000, 800_000),         # crosses the deductible
    (4, date(2026, 2, 3), "GES01", True, False, 3_000_000, 600_000),         # window in 2026: the 130 UF rule
    (4, date(2026, 5, 14), "GES01", True, False, 2_500_000, 500_000),
    (5, date(2025, 2, 10), "GES01", True, True, 10_000_000, 2_000_000),
    (5, date(2025, 4, 22), "GES01", True, True, 9_000_000, 1_800_000),
    (5, date(2025, 7, 30), "GES01", True, True, 7_500_000, 1_500_000),       # crosses the deductible
    (1, date(2025, 12, 1), "GES01", True, False, 1_000_000, 200_000),        # existing beneficiary: late data
]

_INSTRUCCIONES = [
    "How to use this file",
    "",
    "1. Edit the Beneficiarios and Prestaciones sheets: add rows, change values. Keep the header row as it is.",
    "2. Load it:          python -m src.motor ingerir --db data/motor.db --archivo <this file> --formato A",
    "3. Recalculate:      python -m src.motor recalcular-todos --db data/motor.db --tipo GES",
    "4. Check the result: python -m src.motor ver --db data/motor.db --beneficiario 3",
    "",
    "Beneficiarios (loaded first, so a new beneficiary and their prestaciones can go in the same file)",
    "- ID Beneficiario: whole number. A new ID adds a beneficiary.",
    "- An ID that already exists with the same data is skipped; with different data it is rejected (beneficiaries can't be changed).",
    "- Cotización Pactada UF: greater than 0, e.g. 4.5 or 4,5.",
    "- Fecha Derivación CAEC: a date, or empty if not referred.",
    "",
    "Prestaciones (append-only: loaded rows are never updated)",
    "- ID Beneficiario must exist in the database or in the Beneficiarios sheet.",
    "- Fecha Prestación: an Excel date, YYYY-MM-DD, DD-MM-YYYY or DD/MM/YYYY.",
    "- Código GES: e.g. GES01; empty = not GES. Only GES01 has rules in the seed.",
    "- Red GES / Red CAEC: S/N, Si/No, 1/0. Red GES empty = S.",
    "- Arancel, Copago: whole pesos, e.g. 1000000, 1.000.000 or $ 1.000.000. GES copagos above 20% of the arancel are flagged.",
    "- The rules and UF values in the seed cover 2025-01-01 to 2026-12-31.",
    "",
    "Loading the file again is safe: rows already loaded are rejected as duplicates and only new rows go in.",
    "Rows that can't be read are listed with their row number and reason; the rest are loaded.",
]


def escribir_plantilla(path, columnas_prestaciones: dict[str, str]) -> None:
    """Write the example workbook: Instrucciones, Beneficiarios, Prestaciones (headers from the format)."""
    libro = openpyxl.Workbook()
    ayuda = libro.active
    ayuda.title = "Instrucciones"
    for linea in _INSTRUCCIONES:
        ayuda.append([linea])
    ayuda["A1"].font = Font(bold=True)
    ayuda.column_dimensions["A"].width = 130

    def hoja(titulo, encabezado, filas, formatos):
        h = libro.create_sheet(titulo)
        h.append(encabezado)
        for fila in filas:
            h.append(list(fila))
        for celda in h[1]:
            celda.font = Font(bold=True)
        for letra, formato in formatos.items():
            h.column_dimensions[letra].width = 22
            for (celda,) in h.iter_rows(min_row=2, min_col=ord(letra) - 64, max_col=ord(letra) - 64):
                celda.number_format = formato
        h.freeze_panes = "A2"

    def si_no(b):
        return "S" if b else "N"

    campos_b = ("beneficiario_id", "cotizacion_pactada_uf", "fecha_derivacion_caec")
    hoja(extractor.HOJA_BENEFICIARIOS, [extractor.COLUMNAS_BENEFICIARIOS[c] for c in campos_b],
         EJEMPLO_BENEFICIARIOS, {"A": "0", "B": "0.0", "C": "yyyy-mm-dd"})
    campos_p = ("beneficiario_id", "fecha_prestacion", "codigo_ges", "en_red_ges", "en_red_caec",
                "monto_arancel", "monto_copago")
    hoja(extractor.HOJA_PRESTACIONES, [columnas_prestaciones[c] for c in campos_p],
         [(b, d, c, si_no(rg), si_no(rc), a, cp) for b, d, c, rg, rc, a, cp in EJEMPLO_PRESTACIONES],
         {"A": "0", "B": "yyyy-mm-dd", "C": "@", "D": "@", "E": "@", "F": "#,##0", "G": "#,##0"})
    libro.save(path)
